"""Validar archivos, esquemas administrativos y proyectos existentes; sin modificarlos."""
import csv
import re
import sys
from zipfile import BadZipFile
import yaml
from openpyxl import load_workbook
from _estructura import (ROOT, CSV_HEADERS, SECTOR_DIRS, INTEGRATOR_DIRS,
                         FICHA_FIELDS, TAX_HEADERS, VARIABLE_HEADERS,
                         CATALOG_HEADERS, load_projects, folder_name, read_yaml)


def validate():
    errors = []
    def require(path, directory=False):
        valid = path.is_dir() if directory else path.is_file()
        if not valid:
            errors.append(f'Falta {"carpeta" if directory else "archivo"}: {path.relative_to(ROOT)}')
        return valid

    def csv_check(path, expected):
        if require(path):
            try:
                with path.open(encoding='utf-8-sig', newline='') as stream:
                    reader = csv.reader(stream)
                    if next(reader, None) != expected:
                        errors.append(f'Encabezado CSV incorrecto: {path.relative_to(ROOT)}')
                    for row_number, row in enumerate(reader, 2):
                        if len(row) != len(expected):
                            errors.append(f'Número de campos incorrecto: {path.relative_to(ROOT)}:{row_number}')
            except (OSError, UnicodeError, csv.Error) as error:
                errors.append(f'CSV ilegible: {path.name}: {error}')

    master = ROOT / '00_manual_metodologico'
    for name in ['README.md','.gitignore','pyproject.toml','config/proyectos.yml',
                 'scripts/crear_proyecto.py','scripts/validar_estructura.py','scripts/_estructura.py']:
        require(ROOT / name)
    for name in ['manual_metodologico.qmd','criterios_calidad.qmd','politica_uso_IA.qmd','prompts_maestros.qmd']:
        require(master / name)
    csv_check(master / 'taxonomia_maestra.csv', TAX_HEADERS)
    excel = master / 'diccionario_variables.xlsx'
    if require(excel):
        try:
            wb = load_workbook(excel, read_only=True)
            try:
                if wb.sheetnames != ['variables','catalogos','README']:
                    errors.append('Hojas incorrectas en diccionario_variables.xlsx.')
                for name, expected in [('variables',VARIABLE_HEADERS),('catalogos',CATALOG_HEADERS)]:
                    if name in wb and list(next(wb[name].iter_rows(values_only=True), ())) != expected:
                        errors.append(f'Encabezados incorrectos en hoja {name}.')
                if 'README' in wb and not wb['README']['A1'].value:
                    errors.append('README del diccionario vacío.')
            finally:
                wb.close()
        except (OSError, ValueError, KeyError, BadZipFile) as error:
            errors.append(f'Excel ilegible: {error}')
    try:
        projects = load_projects()
    except (OSError, ValueError, yaml.YAMLError) as error:
        errors.append(f'Registro maestro inválido: {error}')
        projects = {}

    def project_check(path, project=None, template=False):
        if not require(path, True):
            return
        dirs = INTEGRATOR_DIRS if project and project['tipo']=='integrador' else SECTOR_DIRS
        require(path / 'README.md')
        for stage in dirs:
            require(path / stage, True)
        governance = path / '00_gobernanza'
        for name, cols in CSV_HEADERS.items():
            csv_check(governance / name, cols)
        ficha_path = governance / 'ficha_proyecto.yml'
        if require(ficha_path):
            try:
                ficha = read_yaml(ficha_path)
                if not isinstance(ficha, dict):
                    raise ValueError('La ficha debe ser un mapa YAML.')
                for key in FICHA_FIELDS:
                    if key not in ficha:
                        errors.append(f'Falta campo {key}: {path.name}')
                for key in ('subsectores','objetivos_especificos','investigadores'):
                    if not isinstance(ficha.get(key), list):
                        errors.append(f'{key} debe ser lista: {path.name}')
                if ficha.get('poblacion_objetivo') != 'mipymes':
                    errors.append(f'Población incoherente: {path.name}')
                if project:
                    expected = {'project_id':project['id'], **{k:project[k] for k in ('slug','sector','nivel_decisional','estado')}}
                    for key, value in expected.items():
                        if ficha.get(key) != value:
                            errors.append(f'Ficha y registro no coinciden en {key}: {path.name}')
                    if path.name != folder_name(project):
                        errors.append(f'Nombre de carpeta incoherente: {path.name}')
                if template and (ficha.get('project_id') is not None or ficha.get('slug') is not None):
                    errors.append('La plantilla contiene un ID o slug de instancia.')
            except (OSError, ValueError, yaml.YAMLError) as error:
                errors.append(f'Ficha inválida: {path.name}: {error}')

    project_check(ROOT / '_template_proyecto', template=True)
    found = set()
    for path in sorted(ROOT.iterdir()):
        if path.is_dir() and re.match(r'^\d{2}_', path.name) and path.name != '00_manual_metodologico':
            ident = int(path.name[:2])
            project = projects.get(ident)
            if project is None:
                errors.append(f'Carpeta no registrada: {path.name}')
            else:
                if ident in found:
                    errors.append(f'ID de carpeta duplicado: {ident}')
                found.add(ident)
                project_check(path, project)
    for ident in (1,22):
        if ident not in found:
            errors.append(f'Falta proyecto inicial obligatorio: {ident:02d}')
    if errors:
        for error in errors:
            print(f'ERROR: {error}', file=sys.stderr)
        print(f'Validación fallida: {len(errors)} error(es).', file=sys.stderr)
        return 1
    print(f'Validación correcta: maestros, plantilla y {len(found)} proyectos existentes; 22 registrados.')
    return 0


if __name__ == '__main__':
    sys.exit(validate())
